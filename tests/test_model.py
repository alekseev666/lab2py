import os

import pytest
from docx import Document
from openpyxl import load_workbook

from model.premises import Premises
from model.room import Room
from model.flat import Flat
from model.house import House
from model.report import create_report_word, create_report_excel


@pytest.fixture
def house():
    h = House("ул. Ленина, 5")
    f = h.add_flat(101)
    f.add_room("спальня", 4, 5)
    f.add_room("кухня", 3, 3)
    f = h.add_flat(102)
    f.add_room("студия", 5, 6)
    return h


# Комната

def test_room_area_volume():
    r = Room("спальня", 4, 5, 2.5)
    assert r.area() == 20
    assert r.volume() == 50


@pytest.mark.parametrize("a, b, c", [(0, 1, 1), (1, -1, 1), (1, 1, 0)])
def test_room_wrong_size(a, b, c):
    with pytest.raises(ValueError):
        Room("комната", a, b, c)


def test_room_setter():
    r = Room("комната", 2, 2, 2)
    r.a = 3
    assert r.area() == 6
    with pytest.raises(ValueError):
        r.b = -5


def test_room_empty_name():
    with pytest.raises(ValueError):
        Room("  ", 1, 1, 1)


def test_room_str_repr():
    r = Room("кухня", 3, 3, 2.5)
    assert str(r) == "комната кухня (9.0 м²)"
    assert repr(r) == "Room(кухня, 3 x 3 x 2.5)"


# Квартира

def test_flat_area_volume_len():
    f = Flat(1)
    f.add_room("спальня", 4, 5)
    f.add_room("кухня", 3, 3)
    assert f.area() == 29
    assert f.volume() == 72.5
    assert len(f) == 2
    assert repr(f) == "квартира #1 (2 комнат)"


def test_flat_ceiling_is_2_5():
    f = Flat(1)
    r = f.add_room("комната", 4, 5)
    assert r.c == 2.5
    assert f.volume() == 50


@pytest.mark.parametrize("number", [0, -1, 1.5, "1", True])
def test_flat_wrong_number(number):
    with pytest.raises(ValueError):
        Flat(number)


# Дом

def test_house_area_volume_len(house):
    assert house.area() == 59
    assert house.volume() == 147.5
    assert len(house) == 2
    assert repr(house) == "House(ул. Ленина, 5, 2 flats)"


def test_house_duplicate_flat(house):
    with pytest.raises(ValueError):
        house.add_flat(101)


def test_house_empty_address():
    with pytest.raises(ValueError):
        House("   ")


def test_heat_power(house):
    # Q = (V * (T_in - T_out) * K) / 860
    assert house.heat_power(21, -5, 1) == pytest.approx(147.5 * 26 * 1 / 860)


def test_heat_power_room_and_flat():
    r = Room("комната", 4, 5, 2.5)
    assert r.heat_power(20, -20, 1.5) == pytest.approx(50 * 40 * 1.5 / 860)
    f = Flat(1)
    f.add_room("комната", 4, 5)
    assert f.heat_power(20, -20, 1.5) == pytest.approx(r.heat_power(20, -20, 1.5))


@pytest.mark.parametrize("t_in, t_out, k", [(20, 20, 1), (10, 20, 1), (20, -5, 0.5), (20, -5, 5)])
def test_heat_power_wrong_params(house, t_in, t_out, k):
    with pytest.raises(ValueError):
        house.heat_power(t_in, t_out, k)


# Абстрактный класс и наследование

def test_premises_is_abstract():
    with pytest.raises(TypeError):
        Premises()


def test_inheritance():
    assert issubclass(Room, Premises)
    assert issubclass(Flat, Premises)
    assert issubclass(House, Premises)


def test_compare_by_area():
    small = Room("маленькая", 2, 2, 2.5)
    big = Room("большая", 3, 3, 2.5)
    assert small < big
    assert Room("а", 2, 3, 2.5) == Room("б", 3, 2, 2.5)
    assert Room("а", 2, 3, 2.5) != None


# Отчёты

def test_report_word(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    path = create_report_word(59.0, 4.4593)
    assert path == os.path.join("reports", "Отчёт.docx")
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Площадь помещения = 59.00 м²" in text
    assert "Q = 4.46 кВт" in text


def test_report_excel(tmp_path, monkeypatch, house):
    monkeypatch.chdir(tmp_path)
    path = create_report_excel(house, 21, -5, 1)
    assert path == os.path.join("reports", "Отчёт.xlsx")
    ws = load_workbook(path).active
    assert ws["B1"].value == "ул. Ленина, 5"
    assert ws["A7"].value == "Итого"
    assert ws["F7"].value == 59


# Демонстрационный дом

def test_demo_house():
    from main import demo_house
    h = demo_house()
    assert len(h) == 3
    assert h.area() == pytest.approx(59 + 20 + 14 + 7.5)
    assert all(r.c == 2.5 for f in h.flats for r in f.rooms)
