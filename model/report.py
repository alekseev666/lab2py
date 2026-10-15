import os

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font

REPORTS_DIR = "reports"

def report_path(name):
    os.makedirs(REPORTS_DIR, exist_ok=True)
    return os.path.join(REPORTS_DIR, name)

def create_report_word(S,Q):
    doc = Document()
    doc.add_heading('Отчёт', level=1)
    doc.add_paragraph(f"Площадь помещения = {S:.2f} м²")
    doc.add_paragraph(f"Q = {Q:.2f} кВт")
    path = report_path("Отчёт.docx")
    doc.save(path)
    return path

def create_report_excel(house, t_in, t_out, k):
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    bold = Font(bold=True)

    ws.append(["Отчёт", house.address])
    ws.append([])
    ws.append(["Квартира", "Комната", "Длина, м", "Ширина, м", "Высота, м",
               "Площадь, м²", "Объём, м³"])
    for f in house.flats:
        for r in f.rooms:
            ws.append([f.number, r.name, r.a, r.b, r.c,
                       round(r.area(), 2), round(r.volume(), 2)])
    ws.append(["Итого", None, None, None, None,
               round(house.area(), 2), round(house.volume(), 2)])
    total_row = ws.max_row

    ws.append([])
    ws.append(["Температура в здании, °C", t_in])
    ws.append(["Температура на улице, °C", t_out])
    ws.append(["Коэффициент рассеивания", k])
    ws.append(["Тепловая мощность, кВт", round(house.heat_power(t_in, t_out, k), 2)])

    for row in (1, 3, total_row, ws.max_row):
        for cell in ws[row]:
            cell.font = bold
    ws.column_dimensions["A"].width = 28
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 14

    path = report_path("Отчёт.xlsx")
    wb.save(path)
    return path
